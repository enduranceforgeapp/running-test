"""Calcula as fórmulas da planilha de regras e grava tests/planilha_valores.json.

O arquivo .xlsx não guarda os valores calculados, então as fórmulas (todas do tipo
=H2*0.3, =(H2*0.2)+H2, =H2-I2, =J2/(F2-1)) são avaliadas aqui. O JSON é a referência
que tests/test_engine.js compara com a tabela da engine.

Uso: python3 -I tests/planilha_valores.py regras/regras_de_controle_de_volume.xlsx
"""
import json
import re
import sys
from pathlib import Path

import openpyxl

REF = re.compile(r"\b([A-Z]+)(\d+)\b")
SAFE = re.compile(r"^[0-9.+\-*/() ]+$")


def evaluate(ws, coord, cache):
    if coord in cache:
        return cache[coord]
    value = ws[coord].value
    if isinstance(value, str) and value.startswith("="):
        expr = REF.sub(lambda m: repr(evaluate(ws, m.group(0), cache)), value[1:])
        if not SAFE.match(expr):
            raise ValueError(f"fórmula inesperada em {coord}: {value}")
        value = eval(expr, {"__builtins__": {}}, {})  # só números e + - * / ( )
    cache[coord] = value
    return value


def main(path):
    ws = openpyxl.load_workbook(path).worksheets[0]
    cache, rows = {}, []
    for r in range(2, ws.max_row + 1):
        if not isinstance(ws[f"F{r}"].value, (int, float)):
            continue  # linhas de pré-requisito
        rows.append({
            "linha": r,
            "nivel": ws[f"A{r}"].value.strip(),
            "distancia": ws[f"B{r}"].value.strip(),
            "velocidade": ws[f"C{r}"].value.strip(),
            "semanas": ws[f"D{r}"].value.strip(),
            "tempo_minimo": ws[f"E{r}"].value.strip(),
            "sessoes_corrida": ws[f"F{r}"].value,
            "sessoes_resistido": ws[f"G{r}"].value,
            "distancia_semanal": evaluate(ws, f"H{r}", cache),
            "longao": evaluate(ws, f"I{r}", cache),
            "vsdl": evaluate(ws, f"J{r}", cache),
            "media_vsdl": evaluate(ws, f"K{r}", cache),
        })
    out = Path(__file__).with_name("planilha_valores.json")
    out.write_text(json.dumps(rows, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")
    print(f"{len(rows)} linhas gravadas em {out}")


if __name__ == "__main__":
    main(sys.argv[1])
